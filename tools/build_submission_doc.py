#!/usr/bin/env python3
"""Build the FlyReflex competition introduction DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "submission" / "FlyReflex_项目介绍.docx"

BLACK = "000000"
NAVY = "173B57"
BLUE = "1F6E8C"
PALE = "EAF3F7"
LIGHT = "F5F7F8"
BORDER = "D9D9D9"
MUTED = RGBColor(0x4B, 0x55, 0x63)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:color"), BORDER)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))
    set_run_font(run, 9, color=MUTED)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, 10.5, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, 10.5)
    else:
        run = p.add_run(text)
        set_run_font(run, 10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        set_run_font(p.add_run(item), 10.2)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, 16 if level == 1 else 12.5, bold=True, color=RGBColor(0, 0, 0))
    return p


def add_table(doc, headers, rows, widths=None, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(header), 9.5, bold=True, color=RGBColor(255, 255, 255))
        if widths:
            cell.width = widths[col]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cell = cells[col]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, PALE if row_index % 2 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if alignments:
                p.alignment = alignments[col]
            set_run_font(p.add_run(str(value)), 9.2)
            if widths:
                cell.width = widths[col]
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Title", 30), ("Subtitle", 14), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)

    title_p_pr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("FlyReflex 项目介绍  openvela AI Contest 2026"), 8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(82)
    title.paragraph_format.space_after = Pt(10)
    remove_paragraph_border(title)
    set_run_font(title.add_run("FlyReflex 项目介绍"), 30, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.paragraph_format.space_after = Pt(28)
    set_run_font(subtitle.add_run("openvela AI Contest 2026 参赛项目"), 14, color=MUTED)

    lead = doc.add_paragraph()
    lead.paragraph_format.space_after = Pt(20)
    lead.paragraph_format.line_spacing = 1.3
    set_run_font(lead.add_run("基于 MaleCNS 连接组的低时延安全反射层"), 18, bold=True, color=RGBColor(0x17, 0x3B, 0x57))

    add_body(doc, "FlyReflex 在 openvela 上实现一个独立于上层 AI 的确定性安全闭环。当合成视觉输入呈现快速逼近时，局部反射引擎将 LPLC2 和 LC4 两类 looming 特征汇入 DNp01 Giant Fiber 聚合节点，并由 Safety Arbiter 把 AI 的 FORWARD 指令覆盖为 ESCAPE。项目以可追溯数据、固定点实现、场景回放、LVGL 仪表盘和实测时延共同证明其可运行性。")

    add_table(
        doc,
        ["交付维度", "已验证结果"],
        [
            ["安全逻辑", "DANGER 时 AI FORWARD 被覆盖为 ESCAPE"],
            ["平台运行", "goldfish ARM64 镜像交叉编译并在 NSH 执行"],
            ["图形能力", "LVGL 仪表盘自动循环 SAFE SLOW DANGER RECOVERY"],
            ["科学溯源", "MaleCNS v1.0 原始响应 查询语句 311 条直连边完整提交"],
            ["工程验证", "10 类回归断言通过 1000 次模拟器基准完成"],
        ],
        [Cm(3.2), Cm(12.4)],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("项目边界"), 10.5, bold=True)
    add_body(doc, "这是 connectome inspired 的工程原型，不是完整果蝇大脑仿真。连接拓扑和相对结构权重来自 MaleCNS；输入编码、泄漏、阈值、动作映射和 P0 AI 指令均明确标记为工程参数。")

    doc.add_page_break()
    add_heading(doc, "一 科学依据")
    add_body(doc, "MaleCNS v1.0 提供完整雄性果蝇中枢神经系统连接组。项目查询了 LPLC2 和 LC4 到 DNp01 Giant Fiber 的全部直接连接，并将原始 API 响应、Cypher 查询、逐边 CSV、聚合 CSV 和可复现脚本纳入仓库。")

    add_table(
        doc,
        ["直接连接", "突触数", "细胞对边数", "运行时相对权重"],
        [
            ["LPLC2 到 DNp01", "4,862", "185", "433 / 1000"],
            ["LC4 到 DNp01", "6,362", "126", "567 / 1000"],
        ],
        [Cm(5.4), Cm(3.1), Cm(3.1), Cm(4.0)],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "生物回路与工程映射", 2)
    add_table(
        doc,
        ["证据类别", "生物或数据事实", "FlyReflex 中的用法"],
        [
            ["CONNECTOME DERIVED", "LPLC2 和 LC4 直接连接 DNp01", "三节点拓扑与 433 比 567 相对权重"],
            ["PAPER DERIVED", "LPLC2 提供 size 分量 LC4 提供 velocity 分量", "双通道 looming 编码语义"],
            ["ENGINEERING PARAMETER", "无生理数值主张", "固定点泄漏 阈值 clamp 与 ESCAPE 映射"],
        ],
        [Cm(3.6), Cm(6.3), Cm(5.7)],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(doc, "重要说明  突触计数只被用于结构性相对比例，不被解释为生理突触效能。DNp01 的选择同时满足 MaleCNS 直连证据和实验文献中 Giant Fiber 汇聚 looming 通道并参与逃逸起飞的功能证据。", bold_lead="重要说明")

    add_heading(doc, "主要参考来源", 2)
    refs = [
        ("MaleCNS v1.0 数据与下载说明", "https://male-cns.janelia.org/download/"),
        ("Berg 等 Cell 2026 MaleCNS 连接组论文", "https://doi.org/10.1016/j.cell.2026.08.015"),
        ("Ache 等 Current Biology 2019 size velocity 编码", "https://doi.org/10.1016/j.cub.2019.01.079"),
        ("Klapoetke 等 Nature 2017 LPLC2 径向运动选择性", "https://doi.org/10.1038/nature24626"),
    ]
    for label, url in refs:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        add_hyperlink(p, label, url)

    doc.add_page_break()
    add_heading(doc, "二 系统实现")
    add_body(doc, "系统把上层 AI 与本地安全回路并行放置，所有执行命令都必须通过 Safety Arbiter。危险判断不依赖网络、云模型或动态内存分配；热路径只使用固定大小结构和整数运算。")

    add_table(
        doc,
        ["阶段", "输入", "处理", "输出"],
        [
            ["Looming Input", "size 与 expansion rate", "确定性场景或未来相机前端", "视觉事件"],
            ["Feature Channels", "视觉事件", "LPLC2 size proxy 与 LC4 positive rate", "节点活动"],
            ["DNp01 GF", "433 与 567 加权输入", "固定点泄漏 积分 clamp 阈值", "SAFE CAUTION DANGER"],
            ["Safety Arbiter", "AI 指令与反射状态", "DANGER 优先 非法输入 fail safe", "FORWARD ESCAPE STOP"],
        ],
        [Cm(3.0), Cm(3.9), Cm(5.4), Cm(3.3)],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "openvela 集成", 2)
    add_bullets(doc, [
        "作为内建 NSH 应用注册命令 flyreflex 并交叉编译到 goldfish ARM64 镜像",
        "使用 CLOCK MONOTONIC 采集 reflex compute 与 event to arbiter 时延",
        "使用 LVGL 和 openvela framebuffer 呈现状态 节点活动 决策与实时延迟",
        "提供 safe slow danger recovery noise 五类确定性输入和 CSV 日志",
        "非法输入直接输出 STOP 和 FAILSAFE 决策",
    ])

    add_heading(doc, "关键危险路径", 2)
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(8)
    for line in (
        "AI command       FORWARD\n",
        "Reflex state     DANGER\n",
        "Reflex command   ESCAPE\n",
        "Final command    ESCAPE\n",
        "Decision         REFLEX OVERRIDE",
    ):
        run = code.add_run(line)
        run.font.name = "Cascadia Mono"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10)

    add_body(doc, "UI 已在模拟器中完整跑过 SAFE 到 SLOW 到 DANGER 到 RECOVERY 循环；DANGER 帧保持 AI FORWARD 可见，同时显示 ESCAPE 和 REFLEX OVERRIDE，恢复阶段重新交还控制。")

    doc.add_page_break()
    add_heading(doc, "三 验证结果")
    add_heading(doc, "功能与构建", 2)
    add_table(
        doc,
        ["验证项", "结果", "证据"],
        [
            ["主机回归", "通过", "安全 慢速 危险 覆盖 阈值 恢复 衰减 确定性 fail safe 噪声"],
            ["主机构建", "通过", "C11 Release warnings as errors"],
            ["openvela 交叉编译", "通过", "builtin 表和 System map 均包含 flyreflex"],
            ["模拟器危险场景", "通过", "frame 6 产生 DANGER ESCAPE REFLEX OVERRIDE"],
            ["LVGL 运行", "通过", "完整自动场景循环及逐帧日志"],
            ["数据溯源校验", "通过", "数据集 边数 突触数和权重均由脚本断言"],
        ],
        [Cm(4.2), Cm(2.2), Cm(9.4)],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "openvela 模拟器 1000 次基准", 2)
    add_table(
        doc,
        ["路径", "Min ns", "Mean ns", "Median ns", "P95 ns", "P99 ns", "Max ns"],
        [
            ["reflex compute", "1,152", "1,400", "1,200", "2,144", "9,776", "17,552"],
            ["end to end", "2,080", "2,565", "2,144", "3,904", "12,848", "57,872"],
        ],
        [Cm(3.2), Cm(2.0), Cm(2.0), Cm(2.2), Cm(1.9), Cm(1.9), Cm(2.2)],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    add_body(doc, "测量从 looming event 已可用时开始，分别在反射引擎完成和 Safety Arbiter 完成时取 CLOCK MONOTONIC 时间戳。每次迭代重置状态，排序和统计在采样结束后执行。这些数字是 openvela goldfish ARM64 模拟器实测，不代表 MCU，也不包含真实 LLM 往返。")

    doc.add_page_break()
    add_heading(doc, "四 演示与提交")
    add_bullets(doc, [
        "在 NSH 依次运行 flyreflex demo safe  flyreflex demo slow  flyreflex demo danger",
        "运行 flyreflex ui 展示 AI FORWARD 在 DANGER 下被 ESCAPE 覆盖",
        "运行 flyreflex bench 1000 展示模拟器实测统计",
        "短暂展示 aggregate edges 和 model params 证明数据与工程参数分离",
        "结尾说明 synthetic input simulator first 与非完整生物仿真的边界",
    ])
    add_body(doc, "当前代码和文档已经按比赛仓迁移结构组织。提交前仍需由参赛者账号完成官方比赛仓推送 合并 PR 导出官方 AI Coding 会话日志 录制不超过五分钟的视频 并提交文档 视频和仓库链接。")

    add_heading(doc, "五 限制与下一步")
    add_body(doc, "P0 使用合成 looming 输入和模拟 AI FORWARD 指令，未接入真实相机、机器人执行器或云端 LLM；时延结论仅覆盖本地反射与仲裁路径。下一阶段可加入光流或事件相机前端，在低功耗 MCU 或 SoC 上复测，并把 STOP ESCAPE 接入真实执行器，同时保持安全闭环不依赖网络。")

    add_heading(doc, "六 交付物索引")
    add_table(
        doc,
        ["交付物", "内容"],
        [
            ["应用源码", "openvela NSH 应用 固定点引擎 Safety Arbiter LVGL UI"],
            ["数据资产", "MaleCNS 原始响应 311 条直接边 聚合统计 查询与参数分类"],
            ["验证资产", "主机回归测试 模拟器命令输出 1000 次基准"],
            ["研究文档", "架构 科学依据 数据溯源 benchmark demo 开发日志"],
            ["复用能力", "flyreflex provenance Skill 与确定性校验脚本"],
        ],
        [Cm(3.5), Cm(12.1)],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    props = doc.core_properties
    props.title = "FlyReflex 项目介绍"
    props.subject = "openvela AI Contest 2026 参赛项目"
    props.author = "FlyReflex Project"
    props.keywords = "openvela, MaleCNS, connectome, safety reflex, LVGL"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
